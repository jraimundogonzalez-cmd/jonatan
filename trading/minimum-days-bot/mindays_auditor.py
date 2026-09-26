#!/usr/bin/env python3
"""
mindays_auditor.py — Auditor independiente de "Minimum Days" (segunda capa de verificación).

Lee la lista de trades de TradingView y recalcula, sin usar nada del Pine:
  día operativo -> P&L neto diario -> MIN DAYS -> ciclos de N días -> FAIL por drawdown.

ENTRADAS ADMITIDAS
  1) Exportación CSV/XLSX de TradingView (Probador de estrategias -> Lista de operaciones -> descargar).
     Dos filas por trade (entrada y salida). Columnas en inglés o español (ver COLS).
  2) Formato compacto "scrape" (id,lado,entrada,salida,señal,pnl[,mae]) con horas yymmddHHMM.

DEFINICIONES (parámetros entre corchetes)
  - Hora de origen: la del gráfico de TradingView [--tz, por defecto Europe/Madrid]. Se interpreta
    con zoneinfo (DST de Madrid incluido) y se convierte a America/New_York (DST de EE. UU. incluido).
  - Día operativo (sesión CME): de 18:00 ET (D-1) a 17:59:59 ET (D). Hora ET >= 18:00 -> sesión D+1.
    [--session-start 18]
  - Cada trade se asigna a la sesión de su SALIDA.
  - P&L diario = suma del P&L NETO de la exportación (TradingView ya descuenta la comisión en
    "Net P&L"/"Profit"). El auditor NO vuelve a restar comisiones al P&L.
    Se trabaja en céntimos enteros: MIN DAY <=> pnl_cents >= target_cents (sin errores de coma flotante).
  - Días de mercado: sesiones del calendario CME_Equity (pandas_market_calendars) cuyo cierre es
    posterior al inicio de la ventana de señales [--window-start 09:30 ET]; más cualquier día con trades.
  - Ciclo: empieza en el día de mercado siguiente al fin del anterior, con MIN=0, pico=equity inicial.
    Cada día de mercado pertenece al ciclo (sea MIN DAY o no). Termina:
      OK   -> al contar el N-ésimo MIN DAY [--need 4]
      FAIL -> si la equity mínima del día <= pico_EOD_del_ciclo - DD [--dd 2000]
    El pico es TRAILING SOBRE CIERRES DE DÍA (EOD). El DD trailing intradía (tipo Apex) NO se puede
    reconstruir desde la exportación (ver --help-dd).
  - Equity mínima del día, dos cotas (--dd-mode):
      realized : solo equity realizada tras cada trade (COTA SUPERIOR de la equity mínima real ->
                 nunca detecta un FAIL que no ocurrió, puede no detectar alguno que sí).
      mae      : equity antes del trade - comisión de entrada - excursión adversa del trade
                 (COTA INFERIOR -> nunca se le escapa un FAIL, puede marcar alguno de más).
      both     : ejecuta las dos; el valor real de FAILs está entre ambas cifras.

SIMULADOR DE CUENTA HIPOTÉTICA (--account)
  Cada ciclo es una cuenta NUEVA [--starting-equity 25000] con 0 MIN DAYS que opera los trades reales,
  sesión a sesión, trade a trade. Termina con:
      OK                -> al conseguir el N-ésimo MIN DAY [--minimum-days 4]. El ciclo termina EN ESE
                           MOMENTO: su P&L es exactamente el acumulado al alcanzar ese MIN DAY.
      FAIL DRAWDOWN     -> equity <= límite de drawdown [--drawdown 2000, --dd-type]
      FAIL DAILY LOSS   -> P&L del día <= -límite diario [--daily-loss, 0 = off]
  Tras alcanzar el objetivo diario [--minimum-day-profit 150] se bloquean nuevas entradas ese día
  (si la exportación trae trades posteriores, se ignoran y se emite WARNING).
  Dos modelos independientes, calculados por separado y nunca mezclados:
      A (realized): drawdown y pérdida diaria solo con P&L REALIZADO (tras cada trade).
      B (mae)     : P&L realizado + excursión adversa de cada trade mientras está abierto.
                    En un FAIL, la cuenta se liquida justo en el límite violado (sin slippage).
  --dd-type: trailing-eod (pico = máximo cierre de día; por defecto), trailing-intraday (pico =
  máxima equity REALIZADA, actualizada tras cada trade) o static (límite fijo = equity inicial - DD).
  Los parámetros son solo del simulador; no representan las reglas de ninguna empresa.
"""
import argparse
import csv
import datetime as dt
import re
import statistics
import sys
import unicodedata
from dataclasses import dataclass, field
from zoneinfo import ZoneInfo

ET = ZoneInfo("America/New_York")

HELP_DD = """
Qué es "Drawdown" / "Adverse excursion" / "Desviación adversa" en la lista de trades de TradingView:
  - Es la EXCURSIÓN ADVERSA MÁXIMA (MAE) de ESE trade: la peor pérdida no realizada que alcanzó la
    posición entre la entrada y la salida, medida desde el precio de entrada con los máximos/mínimos
    de las velas (o de las sub-velas si Bar Magnifier está activo). Importe en USD = puntos x contratos x
    valor del punto.
  - NO es drawdown de equity de la cuenta, ni drawdown desde un máximo, ni lleva hora.
  - Con trades que no se solapan (pyramiding = 0) la equity entre trades es plana, así que:
        equity_min_del_día = min_i ( equity_antes_del_trade_i - comisión_entrada_i - MAE_i )
    Como máximos/mínimos de vela incluyen todos los ticks, MAE nunca infravalora la peor excursión;
    puede sobrevalorarla si el extremo de la vela de salida ocurrió después de la salida.
    -> sirve como COTA INFERIOR de la equity mínima, no como valor exacto.
  - Para un DD trailing intradía (pico que sube con el P&L NO realizado) haría falta saber el orden
    temporal de la excursión favorable y la adversa dentro de cada trade: la exportación no lo tiene.
"""

COLS = {
    "trade": ["trade #", "trade", "numero de operacion", "n.o de operacion", "no de operacion", "n de operacion"],
    "type": ["type", "tipo"],
    "time": ["date/time", "date and time", "fecha y hora", "fecha/hora", "date", "fecha"],
    "signal": ["signal", "senal"],
    "pnl": ["net p&l usd", "net p&l", "profit usd", "profit", "pyg netas usd", "pyg netas", "pyg neto usd",
            "pyg neto", "p&g netas usd", "p&g neto usd", "beneficio neto usd"],
    "mae": ["adverse excursion usd", "adverse excursion", "drawdown usd", "drawdown",
            "desviacion adversa usd", "desviacion adversa", "excursion adversa usd", "excursion adversa"],
    "commission": ["commission usd", "commission", "comision usd", "comision"],
}


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return " ".join(s.lower().replace("_", " ").split())


def find_col(headers, key):
    nh = [norm(h) for h in headers]
    for cand in COLS[key]:          # 1) coincidencia exacta, en orden de preferencia
        if cand in nh:
            return nh.index(cand)
    for cand in COLS[key]:          # 2) prefijo, excluyendo columnas % y acumuladas
        for i, h in enumerate(nh):
            if h.startswith(cand) and "%" not in h and "cum" not in h and "acum" not in h:
                return i
    return None


def to_cents(s, decimal):
    """Texto -> céntimos enteros. decimal: '.' o ','. Admite '−', 'USD', '$', espacios, miles."""
    if s is None:
        return None
    if isinstance(s, (int, float)):
        return round(float(s) * 100)
    t = str(s).strip().replace("−", "-").replace("USD", "").replace("$", "").replace("\xa0", "").replace(" ", "")
    if t in ("", "-"):
        return None
    if decimal == ",":
        t = t.replace(".", "").replace(",", ".")
    else:
        t = t.replace(",", "")
    return round(float(t) * 100)


def detect_decimal(values):
    """',' si algún valor termina en ,d o ,dd (decimal europeo); si no, '.'."""
    for v in values:
        if isinstance(v, str) and re.search(r",\d{1,2}\s*(USD)?$", v.strip().replace("−", "-")):
            return ","
    return "."


def parse_time(s, src_tz):
    if isinstance(s, dt.datetime):
        naive = s
    else:
        t = str(s).strip()
        naive = None
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M",
                    "%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S", "%y%m%d%H%M"):
            try:
                naive = dt.datetime.strptime(t, fmt)
                break
            except ValueError:
                pass
        if naive is None:
            raise ValueError(f"Formato de fecha no reconocido: {t!r}")
    local = naive.replace(tzinfo=src_tz)             # fold=0 ante horas ambiguas (cambio de hora)
    rt = local.astimezone(dt.timezone.utc).astimezone(src_tz)
    if rt.replace(tzinfo=None) != naive:
        raise ValueError(f"Hora inexistente en {src_tz.key}: {naive} (salto de horario de verano)")
    return local.astimezone(ET)


def session_date(t_et, session_start_h=18):
    """Sesión CME: una hora ET >= session_start pertenece al día siguiente."""
    assert t_et.tzinfo is not None
    t = t_et.astimezone(ET)
    return t.date() + dt.timedelta(days=1) if t.hour >= session_start_h else t.date()


@dataclass
class Trade:
    tid: int
    side: str
    entry_et: dt.datetime
    exit_et: dt.datetime
    pnl_c: int                 # P&L NETO en céntimos (tal cual la exportación)
    mae_c: int = None          # excursión adversa en céntimos, negativa o 0
    comm_entry_c: int = 0      # comisión de entrada (solo para la cota MAE)
    signal: str = ""
    session: dt.date = None


# ----------------------------------------------------------------------------- lectura
def read_rows(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        # La exportación de TradingView trae varias hojas (resumen, análisis...): usar la que
        # tenga las columnas de la lista de operaciones, no simplemente la primera.
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            if rows and find_col(rows[0], "type") is not None and find_col(rows[0], "pnl") is not None:
                return rows
        sys.exit(f"{path}: ninguna hoja tiene la lista de operaciones. Hojas: {wb.sheetnames}")
    with open(path, encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        return list(csv.reader(f, dialect))


def load_tradingview(path, src_tz, comm_side_c=0, decimal="auto"):
    rows = read_rows(path)
    hdr, data = rows[0], [r for r in rows[1:] if any(c not in (None, "") for c in r)]
    ix = {k: find_col(hdr, k) for k in COLS}
    for k in ("trade", "type", "time", "pnl"):
        if ix[k] is None:
            sys.exit(f"Falta la columna '{k}'. Cabeceras: {hdr}")
    dec = detect_decimal([r[ix["pnl"]] for r in data]) if decimal == "auto" else decimal
    by = {}
    for r in data:
        tid = int(float(str(r[ix["trade"]]).strip()))
        typ = norm(r[ix["type"]])
        is_exit = typ.startswith(("exit", "salida", "cierre"))
        is_entry = typ.startswith(("entry", "entrada"))
        if not (is_exit or is_entry):
            sys.exit(f"Tipo de fila desconocido: {r[ix['type']]!r}")
        d = by.setdefault(tid, {})
        t = parse_time(r[ix["time"]], src_tz)
        if is_entry:
            d["entry"] = t
            d["side"] = "L" if ("long" in typ or "larg" in typ) else "S"
        else:
            d["exit"] = t
            d["pnl"] = to_cents(r[ix["pnl"]], dec)
            d["sig"] = str(r[ix["signal"]]) if ix["signal"] is not None else ""
            if ix["mae"] is not None:
                m = to_cents(r[ix["mae"]], dec)
                d["mae"] = -abs(m) if m is not None else None
            if ix["commission"] is not None:
                c = to_cents(r[ix["commission"]], dec)
                d["comm_entry"] = (c or 0) // 2
    out = []
    for tid, d in sorted(by.items()):
        if "exit" not in d:
            continue                                 # trade aún abierto: no hay P&L realizado
        size_comm = d.get("comm_entry", comm_side_c)
        out.append(Trade(tid, d.get("side", "?"), d.get("entry", d["exit"]), d["exit"], d["pnl"],
                         d.get("mae"), size_comm, d.get("sig", "")))
    return out, {"decimal": dec, "columns": {k: (hdr[v] if v is not None else None) for k, v in ix.items()}}


def load_scrape(path, src_tz, comm_side_c=0):
    """id,lado,entrada(yymmddHHMM),salida,señal,pnl[,mae]; 'W' en pnl = +160.08 (compacto)."""
    out = []
    with open(path) as f:
        for r in csv.reader(f):
            if not r or not r[0].strip().isdigit():
                continue
            if len(r) == 5:                          # id,lado,entrada,salida,pnl
                tid, side, e, x, p = r
                sig, mae = "", None
            else:
                tid, side, e, x, sig, p = r[:6]
                mae = r[6] if len(r) > 6 else None
            if p == "W":
                p = "160.08"
            if sig == "T":
                sig = "TP"
            out.append(Trade(int(tid), side, parse_time(e, src_tz), parse_time(x, src_tz), to_cents(p, "."),
                             -abs(to_cents(mae, ".")) if mae not in (None, "") else None, comm_side_c, sig))
    return out, {"decimal": ".", "columns": "scrape"}


# ----------------------------------------------------------------------------- calendario
def _easter(y):
    a, b, c = y % 19, y // 100, y % 100
    h = (19 * a + b - b // 4 - (b - (b + 8) // 25 + 1) // 3 + 15) % 30
    l = (32 + 2 * (b % 4) + 2 * (c // 4) - h - c % 4) % 7
    m = (a + 11 * h + 22 * l) // 451
    return dt.date(y, (h + l - 7 * m + 114) // 31, (h + l - 7 * m + 114) % 31 + 1)


def _observed(d):
    return d - dt.timedelta(days=1) if d.weekday() == 5 else d + dt.timedelta(days=1) if d.weekday() == 6 else d


def market_days(first, last, window_start="09:30", extra=()):
    """Sesiones CME_Equity con cierre posterior al inicio de la ventana NY (+ días con trades)."""
    days = set(extra)
    try:
        import pandas_market_calendars as mcal
        sch = mcal.get_calendar("CME_Equity").schedule(str(first), str(last))
        hh, mm = map(int, window_start.split(":"))
        for idx, row in sch.iterrows():
            close_et = row["market_close"].tz_convert("America/New_York")
            if (close_et.hour, close_et.minute) > (hh, mm):
                days.add(idx.date())
        src = "CME_Equity (pandas_market_calendars)"
    except ImportError:
        # Aproximación: lunes-viernes salvo cierres completos (1-ene, Viernes Santo, 25-dic, observados).
        # El resto de festivos CME tiene sesión NY por la mañana.
        closed = set()
        for y in range(first.year - 1, last.year + 2):
            closed |= {_observed(dt.date(y, 1, 1)), _observed(dt.date(y, 12, 25)), _easter(y) - dt.timedelta(days=2)}
        d = first
        while d <= last:
            if d.weekday() < 5 and d not in closed:
                days.add(d)
            d += dt.timedelta(days=1)
        src = "lunes-viernes sin 1-ene/Viernes Santo/25-dic (aprox.; instala pandas_market_calendars)"
    return sorted(x for x in days if first <= x <= last), src


# ----------------------------------------------------------------------------- motor
@dataclass
class Day:
    date: dt.date
    trades: list = field(default_factory=list)
    pnl_c: int = 0
    low_real_c: int = 0      # mínimo de P&L acumulado del día usando solo P&L realizado (<= 0)
    low_mae_c: int = None    # mínimo usando MAE (<= low_real), None si falta MAE en algún trade


def build_days(trades, session_start_h=18):
    days = {}
    for t in sorted(trades, key=lambda t: t.exit_et):
        t.session = session_date(t.exit_et, session_start_h)
        d = days.setdefault(t.session, Day(t.session))
        d.trades.append(t)
    for d in days.values():
        cum, lo_r, lo_m, mae_ok = 0, 0, 0, True
        for t in d.trades:                           # orden de salida; trades no solapados
            if t.mae_c is None:
                mae_ok = False
            else:
                lo_m = min(lo_m, cum - t.comm_entry_c + t.mae_c, cum + t.pnl_c)
            cum += t.pnl_c
            lo_r = min(lo_r, cum)
        d.pnl_c, d.low_real_c = cum, lo_r
        d.low_mae_c = lo_m if mae_ok else None
    return days


def goal_violations(days, target_c):
    """(B) Trades cuya ENTRADA es posterior al momento en que el realizado del día ya era >= objetivo."""
    bad = []
    for d in days.values():
        cum, lock_t = 0, None
        for t in d.trades:
            if lock_t is not None and t.entry_et >= lock_t:
                bad.append((d.date, t.tid, t.entry_et, lock_t))
            cum += t.pnl_c
            if lock_t is None and cum >= target_c:
                lock_t = t.exit_et
    return bad


def run_cycles(mdays, days, target_c, need, dd_c, mode, start_eq_c=0):
    """Devuelve (lista de ciclos, filas de auditoría diaria). mode: 'realized' | 'mae'."""
    cycles, rows = [], []
    eq = start_eq_c
    i = 0
    cyc_id = 0
    while i < len(mdays):
        cyc_id += 1
        eq0, peak, mins, traded, ntr, res = eq, eq, 0, 0, 0, None
        path = []
        j = i
        while j < len(mdays):
            d = days.get(mdays[j])
            pnl = d.pnl_c if d else 0
            if d:
                traded += 1
                ntr += len(d.trades)
                low = d.low_real_c if mode == "realized" else d.low_mae_c
                if low is None:
                    raise ValueError("modo 'mae' sin columna de excursión adversa")
                intraday_min = eq + low
            else:
                intraday_min = eq
            limit = peak - dd_c
            is_min = d is not None and pnl >= target_c
            if d and intraday_min <= limit:
                eq += pnl
                res = "FAIL"
            else:
                eq += pnl
                if is_min:
                    mins += 1
                peak = max(peak, eq)            # pico trailing sobre cierres de día
                if mins >= need:
                    res = "OK"
            rows.append(dict(cycle=cyc_id, date=mdays[j], trades=len(d.trades) if d else 0, pnl_c=pnl,
                             is_min=is_min, eq_close_c=eq, intraday_min_c=intraday_min, peak_c=peak,
                             dd_c=eq - peak, limit_c=limit, mins=mins,
                             status=res or "OPEN"))
            if d:
                path.append(pnl)
            if res:
                break
            j += 1
        end = min(j, len(mdays) - 1)
        cycles.append(dict(id=cyc_id, res=res or "OPEN", start=mdays[i], end=mdays[end], mins=mins,
                           traded=traded, mkt=end - i + 1, cal=(mdays[end] - mdays[i]).days + 1,
                           trades=ntr, pnl_c=eq - eq0, path=path))
        i = end + 1
    return cycles, rows


def summarize(days, mdays, cycles, target_c):
    td = [days[d] for d in mdays if d in days]
    pn = [d.pnl_c for d in td]
    ok = [c for c in cycles if c["res"] == "OK"]
    f = lambda k: [c[k] for c in ok]
    s = dict(
        market_days=len(mdays), traded_days=len(td), trades=sum(len(d.trades) for d in td),
        min_days=sum(p >= target_c for p in pn),
        cat_A=sum(p >= target_c for p in pn), cat_B=sum(0 <= p < target_c for p in pn),
        cat_C=sum(-target_c < p < 0 for p in pn), cat_D=sum(p <= -target_c for p in pn),
        worst_day=min(pn) / 100 if pn else None, net=sum(pn) / 100,
        cycles_ok=len(ok), cycles_fail=sum(c["res"] == "FAIL" for c in cycles),
        cycles_open=sum(c["res"] == "OPEN" for c in cycles),
        ok_with_profit=sum(c["pnl_c"] > 0 for c in ok))
    s["pct_min"] = round(100 * s["min_days"] / s["traded_days"], 1) if td else None
    if ok:
        for k in ("traded", "mkt", "cal", "trades"):
            v = f(k)
            s[f"{k}_to_ok"] = dict(mean=round(statistics.mean(v), 2), median=statistics.median(v), min=min(v), max=max(v))
        s["pnl_ok_mean"] = round(statistics.mean(f("pnl_c")) / 100, 2)
    best = cur = 0
    for p in pn:
        cur = cur + 1 if p >= target_c else 0
        best = max(best, cur)
    s["max_min_streak"] = best
    return s


def write_daily(rows, path):
    with open(path, "w", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["ciclo", "fecha_sesion", "trades", "pnl_neto", "min_day", "equity_cierre", "equity_min_intradia",
                    "pico_ciclo", "dd_ciclo_cierre", "limite_fail", "estado_ciclo", "min_acumulados"])
        c = lambda v: f"{v / 100:.2f}".replace(".", ",")
        for r in rows:
            w.writerow([r["cycle"], r["date"], r["trades"], c(r["pnl_c"]), "SI" if r["is_min"] else "NO",
                        c(r["eq_close_c"]), c(r["intraday_min_c"]), c(r["peak_c"]), c(r["dd_c"]), c(r["limit_c"]),
                        r["status"], r["mins"]])


# ----------------------------------------------------------------------------- simulador de cuenta
MODELS = (("A", "realized", "A · solo realizado"), ("B", "mae", "B · realizado + excursión adversa"))


def rescale_costs(trades, a):
    """Si --commission/--contracts difieren del backtest, reescala P&L y MAE linealmente (con aviso)."""
    same = round(a.commission * 100) == round(a.bt_commission * 100) and a.contracts == a.bt_contracts
    comm_side_c = round(a.commission * 100) * a.contracts
    if same:
        for t in trades:
            if not t.comm_entry_c:
                t.comm_entry_c = comm_side_c
        return same
    k = a.contracts / a.bt_contracts
    bt_rt_c = round(a.bt_commission * 100) * 2 * a.bt_contracts
    for t in trades:
        t.pnl_c = round((t.pnl_c + bt_rt_c) * k) - 2 * comm_side_c
        t.mae_c = None if t.mae_c is None else round(t.mae_c * k)
        t.comm_entry_c = comm_side_c
    return same


def simulate_account(mdays, days, a, model):
    """Cuenta nueva por ciclo, trade a trade. model: 'realized' | 'mae'. Devuelve (ciclos, trades ignorados)."""
    start, dd, dl = round(a.starting_equity * 100), round(a.drawdown * 100), round(a.daily_loss * 100)
    target = round(a.minimum_day_profit * 100)
    cycles, cur, skipped = [], None, 0
    floor = lambda c: (start if a.dd_type == "static" else c["peak"]) - dd
    for date in mdays:
        if cur is None:
            cur = dict(start=date, eq=start, peak=start, mins=0, mkt=0, traded=0, trades=0, path=[],
                       maxdd=0, mincum=0, res="OPEN", reason="FIN DE DATOS")
        c = cur
        c["mkt"] += 1
        c["end"] = date
        d = days.get(date)
        day, goal, traded, ended = 0, False, False, None
        for t in (d.trades if d else []):
            if goal:                                 # daily stop: sin entradas tras el objetivo
                skipped += 1
                continue
            traded = True
            c["trades"] += 1
            if model == "mae":
                if t.mae_c is None:
                    raise ValueError("modelo B sin columna de excursión adversa")
                low_eq = c["eq"] - t.comm_entry_c + t.mae_c
                low_day = day - t.comm_entry_c + t.mae_c
                c["maxdd"] = max(c["maxdd"], c["peak"] - low_eq)
                c["mincum"] = min(c["mincum"], low_eq - start)
                if low_eq <= floor(c):
                    ended, hit = "DRAWDOWN", floor(c)
                elif dl and low_day <= -dl:
                    ended, hit = "DAILY LOSS", c["eq"] - dl - day
                if ended:                            # liquidación justo en el límite violado
                    day += hit - c["eq"]
                    c["eq"] = hit
                    break
            c["eq"] += t.pnl_c
            day += t.pnl_c
            c["maxdd"] = max(c["maxdd"], c["peak"] - c["eq"])
            c["mincum"] = min(c["mincum"], c["eq"] - start)
            if c["eq"] <= floor(c):
                ended = "DRAWDOWN"
            elif dl and day <= -dl:
                ended = "DAILY LOSS"
            if ended:
                break
            if a.dd_type == "trailing-intraday":
                c["peak"] = max(c["peak"], c["eq"])
            if day >= target:
                goal = True
        if traded:
            c["traded"] += 1
            c["path"].append(day)
        if ended:
            c["res"], c["reason"], c["pnl_c"] = "FAIL", ended, c["eq"] - start
            cycles.append(c)
            cur = None
            continue
        if a.dd_type == "trailing-eod":
            c["peak"] = max(c["peak"], c["eq"])
        if traded and day >= target:
            c["mins"] += 1
        if c["mins"] >= a.minimum_days:
            c["res"], c["reason"] = "OK", f"{a.minimum_days} MIN DAYS"
            c["cal"] = (c["end"] - c["start"]).days + 1
            c["pnl_c"] = c["eq"] - start             # beneficio exacto al alcanzar el último MIN DAY
            cycles.append(c)
            cur = None
    if cur is not None:
        cur["pnl_c"] = cur["eq"] - start
        cycles.append(cur)
    for c in cycles:
        p = c["path"]
        c["neg"] = sum(x < 0 for x in p)
        c["pos_below"] = sum(0 <= x < target for x in p)
        c["worst"] = min(p) if p else None
    return cycles, skipped


def account_stats(cycles):
    ok = [c for c in cycles if c["res"] == "OK"]
    fail = [c for c in cycles if c["res"] == "FAIL"]
    q = lambda xs, fn: fn(xs) if xs else None
    op, cal, pnl = [c["traded"] for c in ok], [c["cal"] for c in ok], [c["pnl_c"] for c in ok]
    negs = [c["neg"] for c in ok]
    hist = {str(k): sum(c["traded"] == k for c in ok) for k in range(1, 10)}
    hist["10+"] = sum(c["traded"] >= 10 for c in ok)
    return dict(
        total=len(ok) + len(fail), ok=len(ok), fail=len(fail),
        fail_dd=sum(c["reason"] == "DRAWDOWN" for c in fail), fail_dl=sum(c["reason"] == "DAILY LOSS" for c in fail),
        open=sum(c["res"] == "OPEN" for c in cycles),
        success=100 * len(ok) / (len(ok) + len(fail)) if ok or fail else None,
        op_mean=q(op, statistics.mean), op_med=q(op, statistics.median), op_max=q(op, max),
        mkt_mean=q([c["mkt"] for c in ok], statistics.mean),
        cal_mean=q(cal, statistics.mean), cal_med=q(cal, statistics.median), cal_max=q(cal, max), hist=hist,
        pnl_mean=q(pnl, statistics.mean), pnl_med=q(pnl, statistics.median), pnl_min=q(pnl, min),
        pnl_max=q(pnl, max), ok_neg=sum(x < 0 for x in pnl),
        clean=sum(n == 0 for n in negs), neg1=sum(n == 1 for n in negs), neg2=sum(n == 2 for n in negs),
        neg3=sum(n >= 3 for n in negs), neg_mean=q(negs, statistics.mean),
        pos_below=sum(c["pos_below"] for c in ok),
        worst_ok=q([c["worst"] for c in ok if c["worst"] is not None], min),
        mincum_mean=q([c["mincum"] for c in ok], statistics.mean), mincum_min=q([c["mincum"] for c in ok], min),
        maxdd_ok_mean=q([c["maxdd"] for c in ok], statistics.mean), maxdd_all=q([c["maxdd"] for c in cycles], max))


def usd(c, d=2):
    """céntimos (o None) -> texto con formato español."""
    if c is None:
        return "-"
    return f"{c / 100:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def num(v, d=2):
    if v is None:
        return "-"
    if isinstance(v, int):
        return str(v)
    return f"{v:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def account_report(trades, days, mdays, cal_src, a, same_costs):
    target = round(a.minimum_day_profit * 100)
    has_mae = all(t.mae_c is not None for t in trades)
    models = MODELS if has_mae else MODELS[:1]
    L = []
    w = L.append
    w("# Simulador de cuenta hipotética — Minimum Days\n")
    w(f"Fuente: `{a.file}` · sesiones {mdays[0]} → {mdays[-1]} · {len(trades)} trades · calendario: {cal_src}\n")
    w("## Parámetros del simulador\n\n| Parámetro | Valor |\n|---|---|")
    for k, v in [("Equity inicial", usd(round(a.starting_equity * 100))),
                 ("Drawdown", f"{usd(round(a.drawdown * 100))} ({a.dd_type})"),
                 ("Límite de pérdida diaria", usd(round(a.daily_loss * 100)) if a.daily_loss else "off"),
                 ("Minimum days", a.minimum_days), ("Minimum day profit", usd(target)),
                 ("Comisión / contrato / lado", num(a.commission)), ("Contratos", a.contracts)]:
        w(f"| {k} | {v} |")
    w("\nParámetros del simulador; no representan las reglas de ninguna empresa.\n")

    viol = goal_violations(days, target)
    w("## Validación del daily stop del Pine\n")
    w(f"Entradas posteriores a alcanzar +{usd(target)} realizado en el día: **{len(viol)}**"
      + (" — ⚠ WARNING: el simulador las ignora." if viol else " (OK)"))
    for v in viol[:10]:
        w(f"- {v[0]} · trade {v[1]} · entrada {v[2]:%Y-%m-%d %H:%M} ET · objetivo alcanzado {v[3]:%H:%M} ET")
    w("")
    w("## Qué se puede reconstruir intradía\n")
    w("- Trades no solapados: el P&L intradía = realizado de los trades cerrados + abierto del trade en curso.")
    w("- La exportación da la excursión adversa máxima (MAE) de cada trade, a la resolución del backtest "
      "(velas / bar magnifier), sin hora ni orden respecto a la excursión favorable.")
    w("- La pérdida diaria INTRADÍA y el drawdown INTRADÍA no se pueden reconstruir exactamente. "
      "Se acotan con dos modelos: **A** (solo realizado: cota optimista) y **B** (realizado + MAE: cota "
      "conservadora). El valor real está entre ambos.")
    if not has_mae:
        w("- ⚠ La exportación NO trae excursión adversa: solo se puede calcular el Modelo A.")
    if not same_costs:
        w(f"- ⚠ Comisión/contratos distintos del backtest ({a.bt_commission} / {a.bt_contracts}): P&L y MAE "
          "reescalados linealmente. No equivale a re-ejecutar el Pine (los ticks de TP/SL cambiarían).")
    w("")

    res = {}
    for key, mode, _ in models:
        cyc, skipped = simulate_account(mdays, days, a, mode)
        res[key] = (cyc, account_stats(cyc), skipped)
    S = {k: v[1] for k, v in res.items()}
    hdr = " | ".join(lbl for _, _, lbl in models)
    sep = "---|" * len(models)

    def row(label, fn):
        w(f"| {label} | " + " | ".join(fn(S[k]) for k, _, _ in models) + " |")

    w(f"## Ciclos\n\n| Métrica | {hdr} |\n|---|{sep}")
    row("Total de ciclos cerrados", lambda s: num(s["total"]))
    row("OK", lambda s: num(s["ok"]))
    row("FAIL", lambda s: num(s["fail"]))
    row("  por drawdown", lambda s: num(s["fail_dd"]))
    row("  por pérdida diaria", lambda s: num(s["fail_dl"]))
    row("Ciclo abierto al final de los datos", lambda s: num(s["open"]))
    row("Success rate (OK / cerrados) %", lambda s: num(s["success"], 1))
    row("Días operados a OK: media", lambda s: num(s["op_mean"]))
    row("  mediana", lambda s: num(s["op_med"], 1))
    row("  máximo", lambda s: num(s["op_max"]))
    row("Días de mercado a OK: media", lambda s: num(s["mkt_mean"]))
    row("Días naturales a OK: media", lambda s: num(s["cal_mean"]))
    row("  mediana", lambda s: num(s["cal_med"], 1))
    row("  máximo", lambda s: num(s["cal_max"]))
    w("\n### Días operados hasta OK\n")
    keys = list(S["A"]["hist"])
    w("| Modelo | " + " | ".join(keys) + " |\n|---|" + "---|" * len(keys))
    for k, _, lbl in models:
        w(f"| {lbl} | " + " | ".join(str(S[k]['hist'][x]) for x in keys) + " |")
    w(f"\n## P&L acumulado al alcanzar el {a.minimum_days}.º Minimum Day (ciclos OK)\n\n| Métrica | {hdr} |\n|---|{sep}")
    row("Media", lambda s: usd(s["pnl_mean"] and round(s["pnl_mean"])))
    row("Mediana", lambda s: usd(s["pnl_med"] and round(s["pnl_med"])))
    row("Mínimo", lambda s: usd(s["pnl_min"]))
    row("Máximo", lambda s: usd(s["pnl_max"]))
    row("Ciclos OK que terminan con P&L < 0", lambda s: num(s["ok_neg"]))
    w(f"\n## Camino hasta OK (ciclos OK)\n\n| Métrica | {hdr} |\n|---|{sep}")
    row("Sin ningún día negativo", lambda s: num(s["clean"]))
    row("Con 1 día negativo", lambda s: num(s["neg1"]))
    row("Con 2 días negativos", lambda s: num(s["neg2"]))
    row("Con 3 o más días negativos", lambda s: num(s["neg3"]))
    row("Días negativos por ciclo: media", lambda s: num(s["neg_mean"]))
    row("Días positivos < objetivo (total)", lambda s: num(s["pos_below"]))
    row("Peor día dentro de un ciclo OK", lambda s: usd(s["worst_ok"]))
    row("Mayor pérdida acumulada: media", lambda s: usd(s["mincum_mean"] and round(s["mincum_mean"])))
    row("Mayor pérdida acumulada: peor", lambda s: usd(s["mincum_min"]))
    row("Máx. drawdown por ciclo OK: media", lambda s: usd(s["maxdd_ok_mean"] and round(s["maxdd_ok_mean"])))
    w("")

    td = [days[d] for d in mdays if d in days]
    n = len(td)
    cats = [(f"A ≥ +{usd(target)}", lambda x: x >= target), (f"B 0 a +{usd(target - 1)}", lambda x: 0 <= x < target),
            (f"C −0,01 a −{usd(target - 1)}", lambda x: -target < x < 0), (f"D ≤ −{usd(target)}", lambda x: x <= -target),
            ("E ≤ −500,00", lambda x: x <= -50000), ("F ≤ −1.000,00", lambda x: x <= -100000),
            ("G ≤ −1.500,00", lambda x: x <= -150000), ("H ≤ −2.000,00", lambda x: x <= -200000)]
    w(f"## Distribución diaria ({n} días operados)\n")
    w("E–H son acumulativas y están incluidas en D. La última columna clasifica cada día por su PEOR punto "
      "intradía (realizado + MAE); solo aplica a las categorías de pérdida.\n")
    w("| Categoría | Días (P&L del día) | % | Días (peor punto intradía) | % |\n|---|---|---|---|---|")
    for i, (name, fn) in enumerate(cats):
        x = sum(fn(d.pnl_c) for d in td)
        if i >= 2 and has_mae:
            y = sum(fn(d.low_mae_c) for d in td)
            yy = (str(y), num(100 * y / n, 1))
        else:
            yy = ("-", "-")
        w(f"| {name} | {x} | {num(100 * x / n, 1)} | {yy[0]} | {yy[1]} |")
    w("")

    per_rows = []
    w("## Estabilidad por periodos (cada periodo simulado por separado, cuentas nuevas)\n")
    w("| Periodo | MIN % | OK A/B | FAIL A/B | Media días a OK A/B | Mediana días a OK A/B | "
      "Peor DD A/B | Trades | Winrate % |\n|---|---|---|---|---|---|---|---|---|")
    for spec in a.periods.split(","):
        y0, _, y1 = spec.strip().partition("-")
        p0, p1 = dt.date(int(y0), 1, 1), dt.date(int(y1 or y0), 12, 31)
        pm = [d for d in mdays if p0 <= d <= p1]
        ptd = [days[d] for d in pm if d in days]
        if not ptd:
            w(f"| {spec} | sin datos | | | | | | | |")
            continue
        ps = {k: account_stats(simulate_account(pm, days, a, mode)[0]) for k, mode, _ in models}
        g = lambda k, f: "/".join(f(ps[m]) for m, _, _ in models)
        ptr = [t for d in ptd for t in d.trades]
        mp = 100 * sum(d.pnl_c >= target for d in ptd) / len(ptd)
        wr = 100 * sum(t.pnl_c > 0 for t in ptr) / len(ptr)
        per_rows.append((spec, mp, ps))
        w(f"| {spec} | {num(mp, 1)} | {g('ok', lambda s: num(s['ok']))} | {g('fail', lambda s: num(s['fail']))} | "
          f"{g('', lambda s: num(s['op_mean']))} | {g('', lambda s: num(s['op_med'], 1))} | "
          f"{g('', lambda s: usd(s['maxdd_all']))} | {len(ptr)} | {num(wr, 1)} |")
    w("")

    allt = [t for d in td for t in d.trades]
    gp = sum(t.pnl_c for t in allt if t.pnl_c > 0)
    gl = -sum(t.pnl_c for t in allt if t.pnl_c <= 0)
    minpct = 100 * sum(d.pnl_c >= target for d in td) / n
    w(f"## Resultado final — histórico completo\n\n| Métrica | {hdr} |\n|---|{sep}")
    row("Minimum Days %", lambda s: num(minpct, 1))
    row("Ciclos OK", lambda s: num(s["ok"]))
    row("Ciclos FAIL", lambda s: num(s["fail"]))
    row("Media días operados a OK", lambda s: num(s["op_mean"]))
    row("Mediana", lambda s: num(s["op_med"], 1))
    row("Máximo", lambda s: num(s["op_max"]))
    row("Media días naturales", lambda s: num(s["cal_mean"]))
    row("Mediana", lambda s: num(s["cal_med"], 1))
    row("Máximo", lambda s: num(s["cal_max"]))
    row(f"P&L medio al {a.minimum_days}.º MIN", lambda s: usd(s["pnl_mean"] and round(s["pnl_mean"])))
    row(f"P&L mediano al {a.minimum_days}.º MIN", lambda s: usd(s["pnl_med"] and round(s["pnl_med"])))
    worst_real = min(d.pnl_c for d in td)
    worst_low = min(d.low_mae_c for d in td) if has_mae else None
    w(f"| Peor día | {usd(worst_real)} (realizado)" + (f" | {usd(worst_low)} (peor punto intradía)" if has_mae else "") + " |")
    row("Mayor DD dentro de un ciclo", lambda s: usd(s["maxdd_all"]))
    row("Winrate trades %", lambda s: num(100 * sum(t.pnl_c > 0 for t in allt) / len(allt), 1))
    row("Profit Factor", lambda s: num(gp / gl if gl else None, 3))
    w("")
    w("| Periodo | MIN % | OK A/B | FAIL A/B | Días mediana A/B | P&L mediano a OK A/B |\n|---|---|---|---|---|---|")
    for spec, mp, ps in per_rows:
        g = lambda f: "/".join(f(ps[m]) for m, _, _ in models)
        w(f"| {spec} | {num(mp, 1)} | {g(lambda s: num(s['ok']))} | {g(lambda s: num(s['fail']))} | "
          f"{g(lambda s: num(s['op_med'], 1))} | {g(lambda s: usd(s['pnl_med'] and round(s['pnl_med'])))} |")
    w("")
    if any(v[2] for v in res.values()):
        w("Trades ignorados por el daily stop en la simulación: " +
          " · ".join(f"{k} {res[k][2]}" for k, _, _ in models) + "\n")
    return "\n".join(L), res


def write_account_cycles(cycles, path):
    with open(path, "w", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["ciclo", "inicio", "fin", "resultado", "motivo", "min_days", "dias_operados", "dias_mercado",
                    "dias_naturales", "trades", "pnl_final", "max_dd", "mayor_perdida_acumulada",
                    "dias_negativos", "peor_dia", "dias_positivos_bajo_objetivo", "camino"])
        c = lambda v: "" if v is None else f"{v / 100:.2f}".replace(".", ",")
        for i, x in enumerate(cycles, 1):
            w.writerow([f"{i:03d}", x["start"], x["end"], x["res"], x["reason"], x["mins"], x["traded"], x["mkt"],
                        x.get("cal", ""), x["trades"], c(x["pnl_c"]), c(-x["maxdd"]), c(x["mincum"]), x["neg"],
                        c(x["worst"]), x["pos_below"], " | ".join(f"{p / 100:+.2f}" for p in x["path"])])


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("--format", choices=["tv", "scrape"], default="tv")
    ap.add_argument("--tz", default="Europe/Madrid", help="zona horaria del gráfico en TradingView")
    ap.add_argument("--session-start", type=int, default=18)
    ap.add_argument("--window-start", default="09:30")
    ap.add_argument("--target", "--minimum-day-profit", dest="target", type=float, default=150.0)
    ap.add_argument("--need", "--minimum-days", dest="need", type=int, default=4)
    ap.add_argument("--dd", "--drawdown", dest="dd", type=float, default=2000.0)
    ap.add_argument("--dd-mode", choices=["realized", "mae", "both"], default="both")
    ap.add_argument("--comm-side", type=float, default=0.0, help="comisión de ENTRADA por trade ($) para la cota MAE")
    ap.add_argument("--decimal", choices=["auto", ".", ","], default="auto")
    ap.add_argument("--first", help="primer día (YYYY-MM-DD); por defecto la sesión del primer trade")
    ap.add_argument("--last", help="último día (YYYY-MM-DD); por defecto la sesión del último trade")
    ap.add_argument("--out", help="prefijo para CSV de auditoría diaria y de ciclos")
    ap.add_argument("--help-dd", action="store_true")
    g = ap.add_argument_group("simulador de cuenta (--account)")
    g.add_argument("--account", action="store_true", help="simular cuentas hipotéticas (modelos A y B)")
    g.add_argument("--starting-equity", type=float, default=25000.0)
    g.add_argument("--dd-type", choices=["trailing-eod", "trailing-intraday", "static"], default="trailing-eod")
    g.add_argument("--daily-loss", type=float, default=0.0, help="límite de pérdida diaria ($); 0 = off")
    g.add_argument("--commission", type=float, default=0.74, help="$ por contrato y lado")
    g.add_argument("--contracts", type=int, default=4)
    g.add_argument("--bt-commission", type=float, default=0.74, help="comisión con la que se hizo el backtest")
    g.add_argument("--bt-contracts", type=int, default=4, help="contratos con los que se hizo el backtest")
    g.add_argument("--periods", default="2019-2020,2021-2022,2023-2024,2025-2026")
    g.add_argument("--report", help="escribir el informe Markdown del simulador en este archivo")
    a = ap.parse_args(argv)
    if a.help_dd:
        print(HELP_DD)
        return
    tz = ZoneInfo(a.tz)
    comm = round(a.comm_side * 100)
    trades, meta = (load_scrape if a.format == "scrape" else load_tradingview)(a.file, tz, comm, *([] if a.format == "scrape" else [a.decimal]))
    if a.account:
        a.starting_equity, a.drawdown, a.minimum_days, a.minimum_day_profit = a.starting_equity, a.dd, a.need, a.target
        same_costs = rescale_costs(trades, a)
    days = build_days(trades, a.session_start)
    first = dt.date.fromisoformat(a.first) if a.first else min(days)
    last = dt.date.fromisoformat(a.last) if a.last else max(days)
    mdays, cal_src = market_days(first, last, a.window_start, extra=[d for d in days if first <= d <= last])
    target_c, dd_c = round(a.target * 100), round(a.dd * 100)
    print(f"Trades: {len(trades)} | columnas: {meta['columns']} | decimal '{meta['decimal']}' | calendario: {cal_src}")
    if a.account:
        report, res = account_report(trades, days, mdays, cal_src, a, same_costs)
        print(report)
        if a.report:
            with open(a.report, "w") as f:
                f.write(report + "\n")
        if a.out:
            for k, (cyc, _, _) in res.items():
                write_account_cycles(cyc, f"{a.out}_cuenta_ciclos_{k}.csv")
        return res
    viol = goal_violations(days, target_c)
    print(f"(B) Entradas después de alcanzar el objetivo diario: {len(viol)}")
    for v in viol[:10]:
        print("    ", v)
    has_mae = all(t.mae_c is not None for t in trades)
    modes = ["realized", "mae"] if a.dd_mode == "both" else [a.dd_mode]
    if "mae" in modes and not has_mae:
        print("AVISO: no hay columna de excursión adversa -> solo modo 'realized'")
        modes = ["realized"]
    results = {}
    for m in modes:
        cycles, rows = run_cycles(mdays, days, target_c, a.need, dd_c, m)
        s = summarize(days, mdays, cycles, target_c)
        results[m] = (cycles, rows, s)
        print(f"\n=== DD '{m}' ===")
        for k, v in s.items():
            print(f"  {k}: {v}")
        if a.out:
            write_daily(rows, f"{a.out}_diario_{m}.csv")
            with open(f"{a.out}_ciclos_{m}.csv", "w", newline="") as f:
                w = csv.writer(f, delimiter=";")
                w.writerow(["ciclo", "resultado", "inicio", "fin", "min_days", "dias_operados", "dias_mercado",
                            "dias_naturales", "trades", "pnl_ciclo", "camino_pnl_dias_operados"])
                for c in cycles:
                    w.writerow([c["id"], c["res"], c["start"], c["end"], c["mins"], c["traded"], c["mkt"], c["cal"],
                                c["trades"], f"{c['pnl_c'] / 100:.2f}".replace(".", ","),
                                " | ".join(f"{p / 100:+.2f}" for p in c["path"])])
    return results


if __name__ == "__main__":
    main()
